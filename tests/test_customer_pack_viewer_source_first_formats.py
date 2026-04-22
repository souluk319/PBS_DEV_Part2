from __future__ import annotations

import base64
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

import requests
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from play_book_studio.app.intake_api import ingest_customer_pack
from tests.test_customer_pack_read_boundary import _FakeChunkingModel, _FakeEmbeddingModel, _test_server


def _write_sample_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "운영 점검 절차"
    body = slide.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "클러스터 상태를 점검한다."
    command = body.add_paragraph()
    command.text = "oc get nodes"
    command.level = 1
    presentation.save(path)


def _write_sample_png(path: Path) -> None:
    path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+X6M8AAAAASUVORK5CYII="
        )
    )


def _write_pptx_with_figure(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "운영 화면 예시"
    body = slide.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "점검 화면을 기준으로 상태를 확인한다."
    slide.shapes.add_picture(str(image_path), 1500000, 1800000, width=2200000)
    presentation.save(path)


def test_source_first_pptx_and_hwpx_viewer_routes_land() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        pptx_path = root / "viewer-sample.pptx"
        _write_sample_pptx(pptx_path)
        hwp_rows = [
            {
                "book_slug": "sample-hwpx",
                "book_title": "한글 샘플",
                "heading": "1. 개요",
                "section_level": 1,
                "section_path": ["1. 개요"],
                "anchor": "1-개요",
                "source_url": "/tmp/sample.hwpx",
                "viewer_path": "/playbooks/customer-packs/draft/index.html#1-개요",
                "text": "첫 번째 설명",
            }
        ]

        with (
            patch(
                "play_book_studio.intake.private_corpus.load_sentence_model",
                return_value=_FakeEmbeddingModel(),
            ),
            patch(
                "play_book_studio.ingestion.chunking.load_sentence_model",
                return_value=_FakeChunkingModel(),
            ),
        ):
            pptx_result = ingest_customer_pack(
                root,
                {
                    "source_type": "pptx",
                    "uri": str(pptx_path),
                    "title": "운영 점검 절차",
                    "approval_state": "approved",
                },
            )
            with patch(
                "play_book_studio.intake.normalization.builders.extract_hwp_rows_with_unhwp",
                return_value=hwp_rows,
            ):
                hwpx_result = ingest_customer_pack(
                    root,
                    {
                        "source_type": "hwpx",
                        "file_name": "sample.hwpx",
                        "file_bytes": b"fake-hwpx",
                        "title": "한글 샘플",
                        "approval_state": "approved",
                    },
                )

        with _test_server(root) as (base_url, _store, _answerer):
            pptx_response = requests.get(
                f"{base_url}/playbooks/customer-packs/{pptx_result['draft_id']}/index.html",
                timeout=10,
            )
            hwpx_response = requests.get(
                f"{base_url}/playbooks/customer-packs/{hwpx_result['draft_id']}/index.html",
                timeout=10,
            )

        assert pptx_response.status_code == 200
        assert "운영 점검 절차" in pptx_response.text
        assert "oc get nodes" in pptx_response.text

        assert hwpx_response.status_code == 200
        assert "한글 샘플" in hwpx_response.text
        assert "첫 번째 설명" in hwpx_response.text


def test_customer_pack_pptx_viewer_renders_extracted_figure_assets() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        image_path = root / "viewer-figure.png"
        pptx_path = root / "viewer-figure.pptx"
        _write_sample_png(image_path)
        _write_pptx_with_figure(pptx_path, image_path)

        with (
            patch(
                "play_book_studio.intake.private_corpus.load_sentence_model",
                return_value=_FakeEmbeddingModel(),
            ),
            patch(
                "play_book_studio.ingestion.chunking.load_sentence_model",
                return_value=_FakeChunkingModel(),
            ),
        ):
            pptx_result = ingest_customer_pack(
                root,
                {
                    "source_type": "pptx",
                    "uri": str(pptx_path),
                    "title": "운영 화면 예시",
                    "approval_state": "approved",
                },
            )

        figure_assets = [
            dict(asset)
            for asset in (dict(pptx_result.get("book") or {}).get("figure_assets") or [])
            if isinstance(asset, dict)
        ]
        assert figure_assets
        asset_url = str(figure_assets[0]["asset_url"])

        with _test_server(root) as (base_url, _store, _answerer):
            viewer_response = requests.get(
                f"{base_url}/playbooks/customer-packs/{pptx_result['draft_id']}/index.html",
                timeout=10,
            )
            asset_response = requests.get(f"{base_url}{asset_url}", timeout=10)

        assert viewer_response.status_code == 200
        assert asset_response.status_code == 200
        assert asset_response.headers["Content-Type"].startswith("image/")
        assert "/api/customer-packs/assets?draft_id=" in viewer_response.text
        assert str(figure_assets[0]["asset_ref"]) in viewer_response.text
        assert '<img src="' in viewer_response.text
