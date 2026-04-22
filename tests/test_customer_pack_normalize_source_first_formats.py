from __future__ import annotations

import base64
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from play_book_studio.app.intake_api import ingest_customer_pack
from play_book_studio.config.settings import load_settings

from test_customer_pack_read_boundary import _FakeChunkingModel, _FakeEmbeddingModel


def _write_sample_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "OCP 운영 개요"
    body = slide.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "클러스터 상태를 먼저 확인한다."
    step = body.add_paragraph()
    step.text = "oc get nodes"
    step.level = 1

    table = slide.shapes.add_table(2, 2, 0, 0, 3000000, 1200000).table
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "모드"
    table.cell(1, 1).text = "active"
    presentation.save(path)


def _write_table_title_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title_table = slide.shapes.add_table(1, 1, 0, 0, 5000000, 900000).table
    title_table.cell(0, 0).text = "KOMSCO 지급결제플랫폼 아키텍처 설계서"
    date_box = slide.shapes.add_textbox(0, 1500000, 3000000, 500000)
    date_box.text_frame.paragraphs[0].text = "2025. 07. 25"
    presentation.save(path)


def _write_sample_png(path: Path) -> None:
    path.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+X6M8AAAAASUVORK5CYII="
        )
    )


def _write_figure_pptx(path: Path, image_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "운영 화면 예시"
    body = slide.placeholders[1].text_frame
    body.clear()
    body.paragraphs[0].text = "점검 화면을 기준으로 상태를 확인한다."
    slide.shapes.add_picture(str(image_path), 1500000, 1800000, width=2200000)
    presentation.save(path)


def test_pptx_ingest_uses_native_slide_extract() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        source = root / "sample.pptx"
        _write_sample_pptx(source)

        with (
            patch(
                "play_book_studio.intake.private_corpus.load_sentence_model",
                return_value=_FakeEmbeddingModel(),
            ),
            patch(
                "play_book_studio.ingestion.chunking.load_sentence_model",
                return_value=_FakeChunkingModel(),
            ),
        ):
            result = ingest_customer_pack(
                root,
                {
                    "source_type": "pptx",
                    "uri": str(source),
                    "title": "샘플 덱",
                    "approval_state": "approved",
                },
            )

        book = dict(result.get("book") or {})
        sections = [dict(section) for section in (book.get("sections") or []) if isinstance(section, dict)]
        parser_evidence = dict(book.get("parser_evidence") or {})
        quality_contract = dict(book.get("quality_contract") or {})
        scorecard = dict(book.get("parser_challenger_scorecard") or {})
        stage_strategy = dict(book.get("pipeline_stage_strategy") or {})
        playbook_grade = dict(book.get("playbook_grade") or {})
        blocks = [dict(block) for block in (sections[0].get("document_blocks") or []) if isinstance(block, dict)]
        assert result["status"] == "normalized"
        assert sections
        assert sections[0]["heading"] == "OCP 운영 개요"
        assert "클러스터 상태를 먼저 확인한다." in str(sections[0]["text"])
        assert "oc get nodes" in str(sections[0]["text"])
        assert "[TABLE]" in str(sections[0]["text"])
        assert parser_evidence["truth_owner"] == "canonical_book_json"
        assert parser_evidence["primary_parse_strategy"] == "native_ooxml_first"
        assert quality_contract["document_block_count"] >= 4
        assert scorecard["current_candidate_id"] == "pptx_native_slide_extract"
        assert scorecard["recommended_candidate_id"] == "pptx_native_slide_extract"
        assert scorecard["gold_candidate_count"] >= 1
        assert stage_strategy["parse_stage"] == "native slide parse plus challenger comparison"
        assert playbook_grade["grade"] == "silver"
        assert playbook_grade["current_lane_target_grade"] == "silver"
        block_types = [block["block_type"] for block in blocks]
        assert block_types[0] == "heading"
        assert "paragraph" in block_types
        assert "list_item" in block_types
        assert "table" in block_types
        assert block_types.index("paragraph") < block_types.index("list_item")


def test_hwpx_ingest_uses_structured_hwp_rows_when_available() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        rows = [
            {
                "book_slug": "sample-hwpx",
                "book_title": "한글 샘플",
                "heading": "1. 개요",
                "section_level": 1,
                "section_path": ["1. 개요"],
                "anchor": "1-개요",
                "source_url": "/tmp/sample.hwpx",
                "viewer_path": "/playbooks/customer-packs/draft/index.html#1-개요",
                "text": "첫 번째 설명\n\n[TABLE]\n항목 | 값\n모드 | active\n[/TABLE]",
            }
        ]

        with (
            patch(
                "play_book_studio.intake.private_corpus.load_sentence_model",
                return_value=_FakeEmbeddingModel(),
            ),
            patch(
                "play_book_studio.ingestion.chunking.load_sentence_model",
                return_value=_FakeChunkingModel(),
            ),
            patch(
                "play_book_studio.intake.normalization.builders.extract_hwp_rows_with_unhwp",
                return_value=rows,
            ),
        ):
            result = ingest_customer_pack(
                root,
                {
                    "source_type": "hwpx",
                    "file_name": "sample.hwpx",
                    "file_bytes": b"fake-hwpx",
                    "title": "한글 샘플",
                    "approval_state": "approved",
                },
            )

        book = dict(result.get("book") or {})
        sections = [dict(section) for section in (book.get("sections") or []) if isinstance(section, dict)]
        parser_evidence = dict(book.get("parser_evidence") or {})
        scorecard = dict(book.get("parser_challenger_scorecard") or {})
        stage_strategy = dict(book.get("pipeline_stage_strategy") or {})
        playbook_grade = dict(book.get("playbook_grade") or {})
        blocks = [dict(block) for block in (sections[0].get("document_blocks") or []) if isinstance(block, dict)]
        assert result["status"] == "normalized"
        assert sections
        assert sections[0]["heading"] == "1. 개요"
        assert "첫 번째 설명" in str(sections[0]["text"])
        assert "[TABLE]" in str(sections[0]["text"])
        assert parser_evidence["primary_parse_strategy"] == "structured_hwp_first"
        assert scorecard["current_candidate_id"] == "unhwp_structured_rows"
        assert scorecard["recommended_candidate_id"] == "unhwp_structured_rows"
        assert stage_strategy["current_lane_target_grade"] == "bronze"
        assert playbook_grade["grade"] == "silver"
        assert playbook_grade["current_lane_target_grade"] == "bronze"
        assert [block["block_type"] for block in blocks] == ["heading", "paragraph", "table"]


def test_pptx_prefers_table_title_over_date_text() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        source = root / "table-title.pptx"
        _write_table_title_pptx(source)

        with (
            patch(
                "play_book_studio.intake.private_corpus.load_sentence_model",
                return_value=_FakeEmbeddingModel(),
            ),
            patch(
                "play_book_studio.ingestion.chunking.load_sentence_model",
                return_value=_FakeChunkingModel(),
            ),
        ):
            result = ingest_customer_pack(
                root,
                {
                    "source_type": "pptx",
                    "uri": str(source),
                    "title": "테이블 제목 샘플",
                    "approval_state": "approved",
                },
            )

        book = dict(result.get("book") or {})
        sections = [dict(section) for section in (book.get("sections") or []) if isinstance(section, dict)]
        assert result["status"] == "normalized"
        assert sections
        assert sections[0]["heading"] == "KOMSCO 지급결제플랫폼 아키텍처 설계서"
        assert "2025. 07. 25" in str(sections[0]["text"])


def test_pptx_ingest_extracts_figure_assets_for_viewer() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        image_path = root / "sample.png"
        source = root / "figure-sample.pptx"
        _write_sample_png(image_path)
        _write_figure_pptx(source, image_path)

        with (
            patch(
                "play_book_studio.intake.private_corpus.load_sentence_model",
                return_value=_FakeEmbeddingModel(),
            ),
            patch(
                "play_book_studio.ingestion.chunking.load_sentence_model",
                return_value=_FakeChunkingModel(),
            ),
        ):
            result = ingest_customer_pack(
                root,
                {
                    "source_type": "pptx",
                    "uri": str(source),
                    "title": "운영 화면 예시",
                    "approval_state": "approved",
                },
            )

        book = dict(result.get("book") or {})
        sections = [dict(section) for section in (book.get("sections") or []) if isinstance(section, dict)]
        figure_assets = [dict(asset) for asset in (book.get("figure_assets") or []) if isinstance(asset, dict)]
        blocks = [dict(block) for block in (sections[0].get("document_blocks") or []) if isinstance(block, dict)]
        assert result["status"] == "normalized"
        assert sections
        assert figure_assets
        figure_asset = figure_assets[0]
        asset_ref = str(figure_asset["asset_ref"])
        assert asset_ref.startswith("slide-001-figure-01.")
        assert figure_asset["source_page_or_slide"] == "slide:1"
        assert figure_asset["figure_type"] == "slide_image"
        assert figure_asset["asset_url"].startswith("/api/customer-packs/assets?draft_id=")
        assert f'[FIGURE ref="{asset_ref}"' in str(sections[0]["text"])
        assert any(
            block.get("block_type") == "figure" and block.get("figure_asset_ref") == asset_ref
            for block in blocks
        )
        asset_path = load_settings(root).customer_pack_assets_dir / str(result["draft_id"]) / asset_ref
        assert asset_path.exists()
